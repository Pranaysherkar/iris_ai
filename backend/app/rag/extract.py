"""Document text extraction: local → LlamaParse → Gemini vision fallback."""

from __future__ import annotations

import io
import logging
from functools import lru_cache
from typing import Optional

from app.core.config import settings
from app.rag.mime_types import is_docx_mime, is_image_mime, is_plain_text_mime
from app.rag.normalize import normalize_document_text
from app.rag.types import ExtractionResult

logger = logging.getLogger(__name__)


def _extract_plain_text(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n\n".join(parts)


@lru_cache(maxsize=1)
def _llama_client():
    from llama_cloud import LlamaCloud

    key = (settings.LLAMA_CLOUD_API_KEY or "").strip()
    if not key:
        raise RuntimeError("LLAMA_CLOUD_API_KEY is required for LlamaParse")
    return LlamaCloud(api_key=key)


def _llamaparse_extract(data: bytes, file_name: str) -> ExtractionResult:
    client = _llama_client()
    tier = (settings.LLAMAPARSE_TIER or "agentic").strip()
    # parse() = create + wait + get with expand
    result = client.parsing.parse(
        tier=tier,  # type: ignore[arg-type]
        version="latest",
        upload_file=(file_name, data),
        expand=["markdown", "text"],
        timeout=600.0,
    )
    text = (
        getattr(result, "markdown_full", None)
        or getattr(result, "markdown", None)
        or getattr(result, "text_full", None)
        or getattr(result, "text", None)
        or ""
    )
    if isinstance(text, list):
        text = "\n\n".join(str(x) for x in text if x)
    text = normalize_document_text(str(text))
    if not text:
        raise RuntimeError("LlamaParse returned empty text")
    return ExtractionResult(
        text=text,
        source="llamaparse",
        metadata={"tier": tier, "file_name": file_name},
    )


def _gemini_vision_extract(data: bytes, mime: str, file_name: str) -> ExtractionResult:
    from google import genai
    from google.genai import types

    key = (settings.GEMINI_API_KEY or "").strip()
    if not key:
        raise RuntimeError("GEMINI_API_KEY required for vision fallback")
    client = genai.Client(api_key=key)
    prompt = (
        "Extract all readable text from this document image. "
        "Preserve reading order, headings, and lists. "
        "Return plain text only — no commentary."
    )
    part = types.Part.from_bytes(data=data, mime_type=mime)
    response = client.models.generate_content(
        model=settings.GEMINI_VISION_MODEL,
        contents=[prompt, part],
    )
    text = normalize_document_text(getattr(response, "text", None) or "")
    if not text:
        raise RuntimeError("Gemini vision returned empty text")
    return ExtractionResult(
        text=text,
        source="gemini_vision",
        metadata={"model": settings.GEMINI_VISION_MODEL, "file_name": file_name},
    )


def extract_document_text(
    data: bytes,
    *,
    mime: str,
    file_name: str,
) -> ExtractionResult:
    """
    Route by MIME:
    - plain text → local decode
    - docx → python-docx (LlamaParse if thin)
    - pdf / images / doc → LlamaParse; images fall back to Gemini vision
    """
    if not data:
        raise ValueError("Empty file")

    if is_plain_text_mime(mime):
        text = normalize_document_text(_extract_plain_text(data))
        if not text:
            raise RuntimeError("Text file is empty")
        return ExtractionResult(text=text, source="local_txt", metadata={"file_name": file_name})

    if is_docx_mime(mime) and mime.endswith("document"):
        try:
            text = normalize_document_text(_extract_docx(data))
            # Prefer local if we got a reasonable amount of text
            if len(text) >= 40:
                return ExtractionResult(
                    text=text,
                    source="local_docx",
                    metadata={"file_name": file_name},
                )
        except Exception as exc:
            logger.info("local_docx_failed falling_back_llamaparse err=%s", exc)

    # Primary: LlamaParse for PDF / images / remaining office docs
    llama_err: Optional[Exception] = None
    if (settings.LLAMA_CLOUD_API_KEY or "").strip():
        try:
            return _llamaparse_extract(data, file_name)
        except Exception as exc:
            llama_err = exc
            logger.warning("llamaparse_failed file=%s err=%s", file_name, exc)
    else:
        llama_err = RuntimeError("LLAMA_CLOUD_API_KEY not configured")

    if is_image_mime(mime):
        try:
            return _gemini_vision_extract(data, mime, file_name)
        except Exception as vision_err:
            raise RuntimeError(
                f"Extraction failed (LlamaParse: {llama_err}; vision: {vision_err})"
            ) from vision_err

    # Last resort for docx if LlamaParse failed but local had something thin
    if is_docx_mime(mime):
        try:
            text = normalize_document_text(_extract_docx(data))
            if text:
                return ExtractionResult(
                    text=text,
                    source="local_docx",
                    metadata={"file_name": file_name, "fallback": True},
                )
        except Exception:
            pass

    raise RuntimeError(f"Extraction failed: {llama_err}")
